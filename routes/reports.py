import io
import csv
from datetime import datetime, timedelta
from flask import Blueprint, request, jsonify, send_file
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from routes.auth import admin_required
from firebase.firebase_config import (
    get_all_students,
    get_all_attendance_records,
    get_attendance_by_date
)

reports_bp = Blueprint("reports", __name__)

def filter_records(records, start_date=None, end_date=None, kelas=None, nama=None, status=None):
    """Fungsi pembantu menyaring data absensi berdasarkan kriteria."""
    filtered = []
    nama_lower = nama.strip().lower() if nama else None
    kelas_str = kelas.strip().lower() if kelas else None
    status_str = status.strip().lower() if status else None

    for rec in records:
        tgl = rec.get("tanggal", "")
        if start_date and tgl < start_date:
            continue
        if end_date and tgl > end_date:
            continue
        if kelas_str and kelas_str != "semua" and kelas_str != "":
            if str(rec.get("kelas", "")).strip().lower() != kelas_str:
                continue
        if status_str and status_str != "semua" and status_str != "":
            if str(rec.get("status", "")).strip().lower() != status_str:
                continue
        if nama_lower:
            rec_nama = str(rec.get("nama", "")).lower()
            rec_nis = str(rec.get("nis", "")).lower()
            if nama_lower not in rec_nama and nama_lower not in rec_nis:
                continue

        filtered.append(rec)
    return filtered

# =========================================================================
# ADMIN STATISTICS & CHARTS
# =========================================================================

@reports_bp.route("/api/admin/stats", methods=["GET"])
@admin_required
def get_admin_stats():
    """
    Menghitung metrik ringkasan untuk dashboard admin:
    - Total siswa
    - Hadir, Terlambat, Izin, Sakit, Alpa hari ini
    - Rekap per kelas
    - Tren 7 hari terakhir
    """
    students = get_all_students()
    total_students = len(students)

    today_str = datetime.now().strftime("%Y-%m-%d")
    today_records_map = get_attendance_by_date(today_str) or {}
    today_records = list(today_records_map.values()) if isinstance(today_records_map, dict) else []

    hadir_count = 0
    terlambat_count = 0
    izin_count = 0
    sakit_count = 0
    alpa_count = 0

    attended_student_ids = set()

    for r in today_records:
        st = str(r.get("status", "")).capitalize()
        attended_student_ids.add(r.get("userId"))
        if st == "Hadir":
            hadir_count += 1
        elif st == "Terlambat":
            terlambat_count += 1
        elif st == "Izin":
            izin_count += 1
        elif st == "Sakit":
            sakit_count += 1
        elif st == "Alpa":
            alpa_count += 1

    # Siswa yang belum absen dianggap berpotensi Alpa
    unrecorded = total_students - len(attended_student_ids)
    if unrecorded < 0:
        unrecorded = 0
    alpa_total = alpa_count + unrecorded

    # Rekap per kelas
    kelas_stats = {}
    for s in students:
        k = s.get("kelas", "Lainnya")
        if k not in kelas_stats:
            kelas_stats[k] = {"total": 0, "hadir": 0, "terlambat": 0, "absen": 0}
        kelas_stats[k]["total"] += 1

        # Cek apakah sudah hadir hari ini
        uid = s.get("uid")
        if uid in today_records_map:
            rec = today_records_map[uid]
            st = str(rec.get("status", "")).capitalize()
            if st == "Hadir":
                kelas_stats[k]["hadir"] += 1
            elif st == "Terlambat":
                kelas_stats[k]["terlambat"] += 1
            else:
                kelas_stats[k]["absen"] += 1
        else:
            kelas_stats[k]["absen"] += 1

    # Tren 7 hari terakhir
    trend_labels = []
    trend_hadir = []
    trend_terlambat = []
    all_records = get_all_attendance_records()

    # Index records by date
    records_by_date = {}
    for rec in all_records:
        t = rec.get("tanggal")
        if t:
            if t not in records_by_date:
                records_by_date[t] = []
            records_by_date[t].append(rec)

    for i in range(6, -1, -1):
        d = datetime.now() - timedelta(days=i)
        d_str = d.strftime("%Y-%m-%d")
        d_label = d.strftime("%d/%m")
        trend_labels.append(d_label)

        day_list = records_by_date.get(d_str, [])
        h = sum(1 for x in day_list if str(x.get("status", "")).capitalize() == "Hadir")
        t = sum(1 for x in day_list if str(x.get("status", "")).capitalize() == "Terlambat")
        trend_hadir.append(h)
        trend_terlambat.append(t)

    return jsonify({
        "success": True,
        "date": today_str,
        "total_students": total_students,
        "hadir": hadir_count,
        "terlambat": terlambat_count,
        "izin": izin_count,
        "sakit": sakit_count,
        "alpa": alpa_total,
        "kelas_stats": kelas_stats,
        "trend": {
            "labels": trend_labels,
            "hadir": trend_hadir,
            "terlambat": trend_terlambat
        }
    })

# =========================================================================
# REPORTS DATA & EXPORT
# =========================================================================

@reports_bp.route("/api/reports", methods=["GET"])
@admin_required
def get_reports():
    start_date = request.args.get("tanggal_mulai")
    end_date = request.args.get("tanggal_akhir")
    kelas = request.args.get("kelas")
    nama = request.args.get("nama")
    status = request.args.get("status")

    all_records = get_all_attendance_records()
    filtered = filter_records(all_records, start_date, end_date, kelas, nama, status)

    return jsonify({
        "success": True,
        "count": len(filtered),
        "records": filtered
    })

@reports_bp.route("/api/reports/export", methods=["GET"])
@admin_required
def export_reports():
    format_type = request.args.get("format", "csv").lower()
    start_date = request.args.get("tanggal_mulai")
    end_date = request.args.get("tanggal_akhir")
    kelas = request.args.get("kelas")
    nama = request.args.get("nama")
    status = request.args.get("status")

    all_records = get_all_attendance_records()
    filtered = filter_records(all_records, start_date, end_date, kelas, nama, status)

    timestamp_str = datetime.now().strftime("%Y%m%d_%H%M%S")

    if format_type == "csv":
        # Export ke CSV
        si = io.StringIO()
        cw = csv.writer(si)
        cw.writerow(["No", "NIS", "Nama Siswa", "Kelas", "Tanggal", "Jam", "Status Kehadiran", "Metode"])

        for idx, rec in enumerate(filtered, start=1):
            cw.writerow([
                idx,
                rec.get("nis", "-"),
                rec.get("nama", "-"),
                rec.get("kelas", "-"),
                rec.get("tanggal", "-"),
                rec.get("jam", "-"),
                rec.get("status", "-"),
                rec.get("method", "-")
            ])

        mem = io.BytesIO()
        mem.write(si.getvalue().encode("utf-8-sig"))
        mem.seek(0)
        filename = f"Laporan_Absensi_{timestamp_str}.csv"
        return send_file(mem, mimetype="text/csv", as_attachment=True, download_name=filename)

    elif format_type in ("excel", "xlsx"):
        # Export ke Excel (.xlsx) dengan OpenPyXL
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Laporan Absensi"

        # Judul & Header Laporan
        ws.merge_cells("A1:H1")
        title_cell = ws["A1"]
        title_cell.value = "LAPORAN KEHADIRAN SISWA - ABSENSI-AI"
        title_cell.font = Font(name="Arial", size=14, bold=True, color="1E3A8A")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.merge_cells("A2:H2")
        sub_cell = ws["A2"]
        sub_cell.value = f"Diekspor pada: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        sub_cell.font = Font(name="Arial", size=10, italic=True)
        sub_cell.alignment = Alignment(horizontal="center")

        headers = ["No", "NIS", "Nama Siswa", "Kelas", "Tanggal", "Jam", "Status", "Metode"]
        ws.append([])  # Baris 3 kosong
        ws.append(headers)  # Baris 4 header

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
        border_thin = Border(
            left=Side(style='thin', color='D1D5DB'),
            right=Side(style='thin', color='D1D5DB'),
            top=Side(style='thin', color='D1D5DB'),
            bottom=Side(style='thin', color='D1D5DB')
        )

        for col_num in range(1, 9):
            c = ws.cell(row=4, column=col_num)
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center", vertical="center")

        # Isi data
        for idx, rec in enumerate(filtered, start=1):
            row_idx = 4 + idx
            st = str(rec.get("status", "-")).capitalize()
            ws.append([
                idx,
                rec.get("nis", "-"),
                rec.get("nama", "-"),
                rec.get("kelas", "-"),
                rec.get("tanggal", "-"),
                rec.get("jam", "-"),
                st,
                rec.get("method", "-")
            ])
            for col_num in range(1, 9):
                cell = ws.cell(row=row_idx, column=col_num)
                cell.border = border_thin
                if col_num in (1, 2, 5, 6, 7, 8):
                    cell.alignment = Alignment(horizontal="center")

        # Auto width columns
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

        excel_stream = io.BytesIO()
        wb.save(excel_stream)
        excel_stream.seek(0)
        filename = f"Laporan_Absensi_{timestamp_str}.xlsx"
        return send_file(
            excel_stream,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename
        )

    return jsonify({"success": False, "message": "Format export tidak didukung. Gunakan 'csv' atau 'excel'."}), 400
